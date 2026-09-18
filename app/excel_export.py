"""Downloadable attendance Excel for a school and date."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app import config
from app.domain import attendance_sheet, get_school
from app.database import today_str


def excel_path_for(school_name: str, school_id: int, date: str) -> Path:
    safe = "".join(c if c.isalnum() or c in "-_" else "_" for c in school_name)
    return config.EXPORTS_DIR / f"{safe}_{school_id}_{date}.xlsx"


def write_attendance_excel(school_id: int, date: str | None = None) -> Path:
    config.EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    date = date or today_str()
    school = get_school(school_id) or {"name": "School", "id": school_id}
    path = excel_path_for(school["name"], school_id, date)
    rows = attendance_sheet(school_id, date)

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    header_fill = PatternFill("solid", fgColor="0F3D3E")
    header_font = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
    present_fill = PatternFill("solid", fgColor="D1FAE5")
    absent_fill = PatternFill("solid", fgColor="FEE2E2")
    thin = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    ws.merge_cells("A1:H1")
    ws["A1"] = f"{school['name']} — Attendance {date}"
    ws["A1"].font = Font(bold=True, size=14, color="0F3D3E", name="Calibri")

    headers = [
        "Student ID",
        "Name",
        "Class",
        "Parent phone",
        "Status",
        "Time in",
        "Time out",
        "Source",
    ]
    for col, title in enumerate(headers, 1):
        cell = ws.cell(3, col, title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin

    for i, row in enumerate(rows, 4):
        values = [
            row["student_id"],
            row["name"],
            row["class_name"],
            row["parent_phone"],
            row["status"],
            row.get("time_in_display") or "",
            row.get("time_out_display") or "",
            (row.get("source") or "").title(),
        ]
        fill = present_fill if row["is_present"] else absent_fill
        for col, value in enumerate(values, 1):
            cell = ws.cell(i, col, value)
            cell.fill = fill
            cell.border = thin
            cell.alignment = Alignment(vertical="center")

    widths = [14, 26, 12, 16, 12, 12, 12, 12]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.auto_filter.ref = f"A3:H{3 + len(rows)}"
    ws.freeze_panes = "A4"
    wb.save(path)
    return path
